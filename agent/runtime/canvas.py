"""Canvases: a work the user sees beside the conversation, drawn from a template (specs/canvas.md).

A canvas is a folder ``<name>.canvas`` holding ``canvas.json`` (its type and title), the content its
template renders (slides: ``slides.html``; a document: ``doc.md``) and ``assets/``. Every change to
that content is kept in ``.versions/``, so an earlier one can be looked at and restored. The model writes the content with the
ordinary file tools; the desktop renders it with the template's own code, in the same isolated frame
as a page, and draws pictures and prints of it when asked. ``canvas.json`` is how a canvas is named
everywhere: the folder alone is not a file the desktop opens.
"""

from __future__ import annotations

import hashlib
import html
import json
import re
import shutil
import time
from pathlib import Path

MANIFEST = "canvas.json"
SUFFIX = ".canvas"
SLIDES_FILE = "slides.html"
DOC_FILE = "doc.md"
VERSIONS = ".versions"
MAX_VERSIONS = 30
# The desktop renders slides.html whole or not at all.
MAX_SLIDES_BYTES = 2 * 1024 * 1024

TYPES = {
    "slides": "a 16:9 slide deck: one 1920x1080 slide per <section>, shown as a list and presented "
              "full-screen; exported as PDF or PPTX",
    "doc": "a document written section by section with the doc_* tools (doc.md), shown with its outline; "
           "the user can comment on a passage; exported as Word, PDF or HTML with doc_export",
}
ENTRIES = {"slides": SLIDES_FILE, "doc": DOC_FILE}

SLIDES_GUIDE = """# Slides canvas

The deck is {slides}. Edit it with the file tools; the panel beside the conversation follows every
write. After a change, look at the slides you touched with canvas_look(canvas=..., slide=N)
and fix what overflows or crowds before you say the deck is done. Export with
canvas(action="export", canvas=..., format="pdf").

## Format
- slides.html is a fragment: optional <style> blocks, then one <section class="slide"> per slide, in
  order. No <html>, <head> or <body>.
- A slide is 1920x1080 CSS pixels and never scrolls: what does not fit is cut off. Body text is 40px;
  plan about six short lines, one table or one chart per slide.
- Write the deck in the user's language.

## Layouts: one class next to "slide"
- title: a centred <h1> and a subtitle <p>. For the opening slide.
- section: a divider with one large <h1> (and an optional <p>).
- (none): a content slide: <h1>, then <p>, <ul>/<ol>, <table> or <figure>.
- columns: <h1>, then <div class="columns"> with two or three children side by side.
- image: one full-bleed <img>, an optional <h1> over it and <p class="caption">.
- quote: a <blockquote> and <p class="by">.

## Parts
- <p class="lead"> for a larger opening line; .muted, .small, .accent; <strong>.
- Key numbers: <div class="stats"><div><b>42%</b><span>what it counts</span></div>...</div>
- Charts: inline <svg viewBox="0 0 W H" width=".." height=".."> drawn by hand; use currentColor and
  var(--deck-accent) so they follow the theme.
- Images: save them in {assets} and write <img src="assets/name.png">. Nothing loads from the
  network; a web address shows as a broken image.
- Speaker notes: <aside class="notes">...</aside> inside a slide. Hidden on the slide; kept for
  the presenter and exported to PPTX notes.

## Theme
In a <style> block, set any of these on :root to restyle the whole deck:
--deck-bg, --deck-text, --deck-muted, --deck-accent, --deck-accent-text, --deck-font,
--deck-heading-font. For one slide: <section class="slide" style="--deck-bg:#0f172a;--deck-text:#f8fafc">.

## Rules
- No script is needed. One that is added runs isolated: no network, storage or pop-ups.
- One idea per slide; titles under about 40 characters; prefer a chart or numbers to a paragraph.
"""

DOC_GUIDE = """# Document canvas

The document is {doc}, an ordinary doc_* document: its first heading is the title, the text under
it the lead (section "_lead"), every other heading a section. Load the documents tools with
tool_search(family="documents") and write it with them: doc_outline to see its sections, doc_edit to
write the lead, add sections and write each one. The panel beside the conversation shows it with its
outline and follows every change.

The user comments on a passage from the panel; a comment is written into the document as
<!-- @astra: ... --> after the passage. doc_comments lists the open ones with their section and
nearby text: act on each with doc_edit, then resolve it with doc_edit(action="resolve_comment").
Export with doc_export (docx, pdf or html). Images go in {assets} and are written as
![text](assets/name.png).
"""

GUIDES = {"slides": SLIDES_GUIDE, "doc": DOC_GUIDE}


class CanvasError(ValueError):
    """A path that is not a canvas, or a request the canvas cannot satisfy."""


def canvas_root(path: Path) -> Path:
    """The canvas folder that ``path`` (the folder, its canvas.json, or a file inside it) belongs to."""
    path = path.resolve()
    for candidate in (path, *path.parents):
        if candidate.name.endswith(SUFFIX) and (candidate / MANIFEST).is_file():
            return candidate
    raise CanvasError(f"{path} is not in a canvas: a canvas is a folder named <name>{SUFFIX} with a {MANIFEST}.")


def inside_canvas(root: Path, *parts: str) -> Path:
    """Where `parts` of the canvas really are; refused when a link leads out of the canvas folder."""
    root = root.resolve()
    path = root.joinpath(*parts).resolve()
    if path != root and root not in path.parents:
        raise CanvasError(f"{'/'.join(parts)} in {root} leads out of the canvas folder.")
    return path


def read_manifest(root: Path) -> dict:
    try:
        manifest = json.loads((root / MANIFEST).read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise CanvasError(f"{root / MANIFEST} cannot be read: {exc}") from None
    if not isinstance(manifest, dict) or manifest.get("type") not in TYPES:
        raise CanvasError(f"{root / MANIFEST} names no known canvas type ({', '.join(TYPES)}).")
    return manifest


def _slug(title: str) -> str:
    slug = re.sub(r"[^\w一-鿿-]+", "-", title.strip(), flags=re.UNICODE).strip("-._")
    return slug[:40] or "canvas"


def create(workdir: Path, kind: str, title: str, path: str = "") -> Path:
    """Make a new canvas of ``kind`` inside ``workdir``; return its folder."""
    if kind not in TYPES:
        raise CanvasError(f"Unknown canvas type {kind!r}; types: {', '.join(TYPES)}.")
    title = title.strip() or "Untitled"
    workdir = workdir.resolve()
    if path.strip():
        target = Path(path.strip()).expanduser()
        target = (target if target.is_absolute() else workdir / target).resolve()
        if not target.name.endswith(SUFFIX):
            target = target.with_name(target.name + SUFFIX)
        if target.exists():
            raise CanvasError(f"{target} already exists; edit it, or choose another path.")
    else:
        stem, number = _slug(title), 1
        target = workdir / f"{stem}{SUFFIX}"
        while target.exists():
            number += 1
            target = workdir / f"{stem}-{number}{SUFFIX}"
    if target != workdir and workdir not in target.parents:
        raise CanvasError("A canvas is made inside the workspace.")
    (target / "assets").mkdir(parents=True)
    manifest = {"type": kind, "title": title, "version": 1, "entry": ENTRIES[kind],
                "created": time.strftime("%Y-%m-%dT%H:%M:%S%z")}
    # One line ending on every platform: the desktop and the doc_* tools read and compare these files as written.
    (target / MANIFEST).write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    (target / ENTRIES[kind]).write_text(starter_slides(title) if kind == "slides" else f"# {title}\n\n",
                                        encoding="utf-8", newline="\n")
    record_version(target, "create")
    return target


def content_file(root: Path) -> Path:
    """The file a canvas's template renders: slides.html or doc.md, where it really is inside the canvas."""
    return inside_canvas(root, ENTRIES[read_manifest(root)["type"]])


def list_versions(root: Path) -> list[dict]:
    """The kept versions of a canvas's content, oldest first."""
    versions = []
    try:
        kept = inside_canvas(root, VERSIONS)
    except CanvasError:
        return versions
    for folder in sorted(kept.glob("[0-9][0-9][0-9][0-9]")):
        try:
            meta = json.loads(inside_canvas(root, VERSIONS, folder.name, "meta.json").read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError, CanvasError):
            continue
        if isinstance(meta, dict):
            versions.append({**meta, "id": folder.name})
    return versions


def record_version(root: Path, source: str) -> str | None:
    """Keep the canvas's content as a new version, unless it is the same as the last one; drop the oldest."""
    content = content_file(root)
    data = content.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    versions = list_versions(root)
    if versions and versions[-1].get("sha256") == digest:
        return None
    kept = inside_canvas(root, VERSIONS)
    kept.mkdir(parents=True, exist_ok=True)
    # Numbered past every folder there, a version that cannot be read included.
    taken = [int(folder.name) for folder in kept.glob("[0-9][0-9][0-9][0-9]")]
    number = max(taken, default=0) + 1
    while True:
        if number > 9999:
            raise CanvasError("This canvas has used up its version numbers.")
        folder = kept / f"{number:04d}"
        try:
            folder.mkdir()
            break
        except FileExistsError:
            number += 1
    folder = inside_canvas(root, VERSIONS, folder.name)
    (folder / content.name).write_bytes(data)
    (folder / "meta.tmp").write_text(json.dumps({"at": time.time(), "source": source, "sha256": digest,
                                                  "file": content.name}, ensure_ascii=False), encoding="utf-8")
    (folder / "meta.tmp").replace(folder / "meta.json")
    for old in versions[: max(0, len(versions) + 1 - MAX_VERSIONS)]:
        shutil.rmtree(kept / old["id"], ignore_errors=True)
    return folder.name


def starter_slides(title: str) -> str:
    safe = title.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return f'<section class="slide title">\n  <h1>{safe}</h1>\n  <p class="muted">{time.strftime("%Y-%m-%d")}</p>\n</section>\n'


_SLIDE = re.compile(r"<section\b[^>]*\bclass\s*=\s*(['\"])(?:[^'\"]*\s)?slide(?:\s[^'\"]*)?\1", re.IGNORECASE)


def _text(markup: str) -> str:
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", markup))).strip()


def slide_parts(root: Path) -> list[tuple[str, str]]:
    """Each slide's first heading and its speaker notes, as plain text, in order."""
    try:
        source = inside_canvas(root, SLIDES_FILE).read_text(encoding="utf-8")
    except (OSError, CanvasError):
        return []
    starts = [match.start() for match in _SLIDE.finditer(source)]
    parts = []
    for index, start in enumerate(starts):
        chunk = source[start: starts[index + 1] if index + 1 < len(starts) else len(source)]
        notes = re.findall(r"<aside\b[^>]*\bclass\s*=\s*['\"][^'\"]*\bnotes\b[^'\"]*['\"][^>]*>(.*?)</aside>", chunk, re.IGNORECASE | re.DOTALL)
        heading = re.search(r"<h[1-3]\b[^>]*>(.*?)</h[1-3]>", re.sub(r"<aside\b.*?</aside>", "", chunk, flags=re.IGNORECASE | re.DOTALL), re.IGNORECASE | re.DOTALL)
        parts.append((_text(heading.group(1)) if heading else "", "\n".join(_text(note) for note in notes)))
    return parts


def build_pptx(pictures: list[Path], parts: list[tuple[str, str]], target: Path, title: str) -> None:
    """A 16:9 deck of one full-slide picture each, the slide's heading as the picture's text for screen readers, its notes as notes."""
    from pptx import Presentation
    from pptx.util import Emu

    deck = Presentation()
    deck.slide_width, deck.slide_height = Emu(12192000), Emu(6858000)
    width, height = deck.slide_width, deck.slide_height
    blank = deck.slide_layouts[6]
    for index, picture_path in enumerate(pictures):
        slide = deck.slides.add_slide(blank)
        picture = slide.shapes.add_picture(str(picture_path), Emu(0), Emu(0), width=width, height=height)
        heading, notes = parts[index] if index < len(parts) else ("", "")
        picture.name = f"Slide {index + 1}"
        # The picture's description is what a screen reader says for the slide.
        properties = next(node for node in picture._element.iter() if str(node.tag).endswith("}cNvPr"))
        properties.set("descr", heading or f"Slide {index + 1}")
        frame = slide.notes_slide.notes_text_frame if notes else None
        if frame is not None:
            frame.text = notes
    deck.core_properties.title = title
    target.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(target))


def slide_count(root: Path) -> int:
    """How many slides slides.html holds."""
    try:
        source = inside_canvas(root, SLIDES_FILE).read_text(encoding="utf-8")
    except (OSError, CanvasError):
        return 0
    return len(_SLIDE.findall(source))
