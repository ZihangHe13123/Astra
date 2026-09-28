"""Export Markdown documents to Word, PDF or HTML.

Two engines produce Word and HTML files:

- ``python-docx``: built in through the optional ``documents`` extra
  (python-docx + markdown-it-py).
- ``pandoc``: used when it is installed (``PATH`` or ``ASTRA_PANDOC``);
  ``auto`` prefers it because it covers more Markdown.

PDF is always a Word file converted by LibreOffice (``soffice``), which keeps
both engines' output identical in layout. Neither engine fetches remote images,
and local images must lie inside the filesystem policy's allowed roots.
"""

from __future__ import annotations

import importlib.util
import os
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.parse
import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .documents import _MARKER, fenced_lines, heading_of, parse_document, remove_comment

FORMATS = ("docx", "pdf", "html")
ENGINES = ("auto", "python-docx", "pandoc")
SUFFIXES = {"docx": ".docx", "pdf": ".pdf", "html": ".html"}
PICTURE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}
CONVERTIBLE_SUFFIXES = {".webp"}
INSTALL_HINT = (
    "Install the Word exporter with `astra setup --extra documents` (in a source checkout: "
    "`uv sync --inexact --extra documents`), then restart Astra; or install pandoc."
)
_IMAGE_REF = re.compile(r"(!\[(?P<alt>[^\]]*)\]\()(?P<src><[^>]*>|[^)\s]+)(?P<rest>(?:\s+\"[^\"]*\")?\))")
_REMOTE = re.compile(r"^[a-z][a-z0-9+.-]*:", re.IGNORECASE)

ImageResolver = Callable[[str], Path]


class ExportError(RuntimeError):
    def __init__(self, code: str, message: str, hint: str = ""):
        super().__init__(message)
        self.code = code
        self.hint = hint


@dataclass
class ExportResult:
    output: Path
    engine: str
    warnings: list[str] = field(default_factory=list)


def python_docx_available() -> bool:
    return bool(importlib.util.find_spec("docx") and importlib.util.find_spec("markdown_it"))


def markdown_it_available() -> bool:
    return bool(importlib.util.find_spec("markdown_it"))


def _configured_program(variable: str) -> str | None:
    configured = os.getenv(variable, "").strip()
    if not configured:
        return None
    found = shutil.which(configured)
    if found:
        return found
    path = Path(configured).expanduser()
    return str(path) if path.is_file() else None


def find_pandoc() -> str | None:
    return _configured_program("ASTRA_PANDOC") or shutil.which("pandoc")


def find_soffice() -> str | None:
    configured = _configured_program("ASTRA_SOFFICE")
    if configured:
        return configured
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    candidates: list[Path] = []
    if sys.platform == "darwin":
        candidates.append(Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"))
    elif os.name == "nt":
        for base in (os.getenv("ProgramFiles", r"C:\Program Files"), os.getenv("ProgramFiles(x86)", r"C:\Program Files (x86)")):
            candidates.append(Path(base) / "LibreOffice" / "program" / "soffice.exe")
    return next((str(item) for item in candidates if item.is_file()), None)


def prepare_markdown(text: str) -> tuple[str, int, str | None]:
    """Strip section markers and review comments; return (markdown, pending, title)."""
    document = parse_document(text)
    pending = sum(item.pristine_pending for item in document.sections)
    title = document.title
    while document.comments:
        document = parse_document("".join(remove_comment(document, document.comments[0].id)))
    fenced = fenced_lines(document.lines)
    kept = [line for index, line in enumerate(document.lines) if fenced[index] or not _MARKER.match(line)]
    return "".join(kept), pending, title


def _clean_src(raw: str) -> str:
    src = raw.strip()
    if src.startswith("<") and src.endswith(">"):
        src = src[1:-1]
    return src


def check_images(markdown: str, resolve_image: ImageResolver, warnings: list[str]) -> tuple[str, dict[str, Path]]:
    """Resolve every image reference; replace unusable ones with their alt text.

    Returns the rewritten Markdown and a map from each kept reference to its file.
    """
    lines = markdown.splitlines(keepends=True)
    fenced = fenced_lines(lines)
    resolved: dict[str, Path] = {}

    def replace(match: re.Match[str]) -> str:
        src = _clean_src(match.group("src"))
        alt = match.group("alt")
        label = f"[image: {alt or src}]"
        if _REMOTE.match(src) and not src.lower().startswith("file:"):
            warnings.append(f"Remote image not embedded (download it first): {src}")
            return label
        try:
            path = resolve_image(src)
        except (OSError, ValueError, PermissionError) as exc:
            warnings.append(f"Image skipped: {src} ({exc})")
            return label
        if not path.is_file():
            warnings.append(f"Image not found: {src}")
            return label
        suffix = path.suffix.lower()
        if suffix not in PICTURE_SUFFIXES | CONVERTIBLE_SUFFIXES:
            warnings.append(f"Image format {suffix or 'unknown'} is not supported in Word; use PNG or JPEG: {src}")
            return label
        resolved[src] = path
        return match.group(0)

    output = [line if fenced[index] else _IMAGE_REF.sub(replace, line) for index, line in enumerate(lines)]
    return "".join(output), resolved


def _lookup(images: dict[str, Path], src: str) -> Path | None:
    """Find a checked image; markdown-it percent-encodes non-ASCII paths."""
    return images.get(src) or images.get(urllib.parse.unquote(src))


def _atomic_move(source: Path, target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_name(f".{target.name}.agent-{os.getpid()}.tmp")
    try:
        shutil.copyfile(source, temporary)
        with open(temporary, "rb") as handle:
            os.fsync(handle.fileno())
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)


def _run(args: list[str], *, timeout: int, cwd: Path | None = None) -> subprocess.CompletedProcess[str]:
    try:
        return subprocess.run(
            args, capture_output=True, text=True, timeout=timeout, cwd=cwd, check=False,
            stdin=subprocess.DEVNULL,
        )
    except subprocess.TimeoutExpired as exc:
        raise ExportError("export_timeout", f"{Path(args[0]).name} did not finish within {timeout} s") from exc
    except OSError as exc:
        raise ExportError("export_unavailable", f"Could not start {args[0]}: {exc}") from exc


# ---- pandoc ------------------------------------------------------------------------------


def export_with_pandoc(
    pandoc: str,
    markdown: str,
    *,
    base_dir: Path,
    output: Path,
    fmt: str,
    title: str | None,
    template: Path | None,
    warnings: list[str],
    paper: str = "a4",
) -> None:
    with tempfile.TemporaryDirectory(prefix="astra-export-") as work:
        source = Path(work) / "input.md"
        source.write_text(markdown, encoding="utf-8")
        target = Path(work) / f"output{SUFFIXES['html' if fmt == 'html' else 'docx']}"
        args = [
            pandoc, str(source), "--from", "gfm", "--to", "html5" if fmt == "html" else "docx",
            "--output", str(target), f"--resource-path={base_dir}",
        ]
        if fmt == "html":
            args += ["--standalone", "--metadata", f"title={title or output.stem}"]
        if fmt != "html" and _title_mode(markdown):
            # Like the built-in exporter: the only level-1 heading becomes the Title and
            # the other headings move up a level. pandoc takes the Title only from a
            # heading that opens the document and would turn any other one into text.
            args.append("--shift-heading-level-by=-1")
        if template is not None and fmt != "html":
            args += ["--reference-doc", str(template)]
        completed = _run(args, timeout=180, cwd=base_dir)
        if completed.returncode != 0 or not target.is_file():
            detail = (completed.stderr or completed.stdout or "").strip()[-2000:]
            raise ExportError("export_failed", f"pandoc failed ({completed.returncode}): {detail}")
        warnings.extend(line.strip() for line in completed.stderr.splitlines()[:10] if line.strip())
        if fmt != "html" and template is None and zipfile.is_zipfile(target):
            _polish_pandoc_docx(target, paper)
        _atomic_move(target, output)


def _title_mode(markdown: str) -> bool:
    """Whether the document opens with its only level-1 heading."""
    lines = markdown.splitlines(keepends=True)
    fenced = fenced_lines(lines)
    headings = [(index, found[0]) for index, line in enumerate(lines)
                if not fenced[index] and (found := heading_of(line.rstrip("\r\n"))) is not None]
    if not headings or headings[0][1] != 1 or sum(level == 1 for _, level in headings) != 1:
        return False
    return not "".join(lines[:headings[0][0]]).strip()


_PAGE_TWIPS = {"a4": (11906, 16838), "letter": (12240, 15840)}
_MARGIN_TWIPS = 1417  # 25 mm, as in the built-in exporter
_GRID_BORDERS = "<w:tblBorders>" + "".join(
    f'<w:{side} w:val="single" w:sz="4" w:space="0" w:color="auto" />'
    for side in ("top", "left", "bottom", "right", "insideH", "insideV")
) + "</w:tblBorders>"


def _polish_pandoc_docx(docx: Path, paper: str) -> None:
    """Give pandoc's default reference layout the built-in exporter's page and tables.

    pandoc leaves the page size to the reader (Letter in LibreOffice), sizes tables to its
    own narrower text width and draws only a header rule. Use the requested paper with
    25 mm margins, full-width tables and grid borders, as the python-docx exporter does.
    """
    width, height = _PAGE_TWIPS.get(paper, _PAGE_TWIPS["a4"])
    text_width = width - 2 * _MARGIN_TWIPS
    with zipfile.ZipFile(docx) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    document = next((data for info, data in entries if info.filename == "word/document.xml"), None)
    if document is None:
        return
    xml = document.decode("utf-8")

    def table_properties(match: re.Match[str]) -> str:
        properties = re.sub(r"<w:tblW\b[^>]*/>", '<w:tblW w:type="pct" w:w="5000" />', match.group(0))
        if "<w:tblBorders" not in properties:
            # Schema order: tblW comes before tblBorders.
            properties = re.sub(r"(<w:tblW\b[^>]*/>)", lambda found: found.group(1) + _GRID_BORDERS, properties, count=1)
        return properties

    def table_grid(match: re.Match[str]) -> str:
        columns = [int(value) for value in re.findall(r'<w:gridCol w:w="(\d+)"', match.group(0))]
        total = sum(columns)
        if not total:
            return match.group(0)
        return "<w:tblGrid>" + "".join(
            f'<w:gridCol w:w="{max(1, round(text_width * value / total))}" />' for value in columns
        ) + "</w:tblGrid>"

    xml = re.sub(r"<w:tblPr>.*?</w:tblPr>", table_properties, xml, flags=re.S)
    xml = re.sub(r"<w:tblGrid>.*?</w:tblGrid>", table_grid, xml, flags=re.S)
    if "<w:pgSz" not in xml:
        margin = _MARGIN_TWIPS
        page = (f'<w:pgSz w:w="{width}" w:h="{height}" /><w:pgMar w:top="{margin}" w:right="{margin}" '
                f'w:bottom="{margin}" w:left="{margin}" w:header="708" w:footer="708" w:gutter="0" />')
        if re.search(r"<w:sectPr\b[^>]*/>", xml):
            xml = re.sub(r"<w:sectPr\b([^>]*)/>", lambda found: f"<w:sectPr{found.group(1)}>{page}</w:sectPr>", xml, count=1)
        elif "</w:sectPr>" in xml:
            # pandoc's section only holds footnote settings, which precede pgSz in the schema.
            xml = xml.replace("</w:sectPr>", page + "</w:sectPr>", 1)
        else:
            xml = xml.replace("</w:body>", f"<w:sectPr>{page}</w:sectPr></w:body>", 1)
    with zipfile.ZipFile(docx, "w", zipfile.ZIP_DEFLATED) as archive:
        for info, data in entries:
            archive.writestr(info, xml.encode("utf-8") if info.filename == "word/document.xml" else data)


# ---- LibreOffice -------------------------------------------------------------------------


_MAC_FONT_DIRS = ("/System/Library/Fonts", "/System/Library/Fonts/Supplemental", "/Library/Fonts", "~/Library/Fonts")
_MAC_CJK_FONTS = (
    ("Hiragino Sans GB", "/System/Library/Fonts/Hiragino Sans GB.ttc"),
    ("STHeiti", "/System/Library/Fonts/STHeiti Light.ttc"),
    ("Songti SC", "/System/Library/Fonts/Supplemental/Songti.ttc"),
    ("Arial Unicode MS", "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
)
_CJK_TEXT = re.compile(r"[\u3040-\u30ff\u3400-\u9fff\uf900-\ufaff\uac00-\ud7af]")


def _link_system_fonts(profile: Path) -> None:
    """Expose macOS system fonts through the throwaway LibreOffice profile.

    Recent LibreOffice builds on macOS can fail to see system fonts when run
    headless, silently substituting bundled fonts and dropping CJK glyphs.
    Fonts under ``<profile>/user/fonts`` are always loaded.
    """
    if sys.platform != "darwin":
        return
    fonts = profile / "user" / "fonts"
    fonts.mkdir(parents=True, exist_ok=True)
    for index, raw in enumerate(_MAC_FONT_DIRS):
        source = Path(raw).expanduser()
        if source.is_dir():
            try:
                (fonts / f"system-{index}").symlink_to(source, target_is_directory=True)
            except OSError:
                continue


# Office fonts a stock Mac lacks and LibreOffice has no stand-in for (Calibri and Cambria
# get Carlito and Caladea): family, font file name prefix, installed replacement.
_MAC_FONT_STANDINS = (
    ("Aptos", "aptos", "Helvetica Neue"),
    ("Aptos Display", "aptos", "Helvetica Neue"),
    ("Consolas", "consola", "Menlo"),
)


def _replace_missing_fonts(profile: Path) -> None:
    """Map Office fonts this Mac lacks to installed ones in the throwaway profile.

    Otherwise LibreOffice falls back to an unrelated font (a rounded Japanese one on
    macOS), which hits pandoc's Aptos body text and the Consolas code font. LibreOffice
    only applies replacement entries marked Always, so installed families are skipped.
    """
    if sys.platform != "darwin":
        return
    installed: set[str] = set()
    for raw in _MAC_FONT_DIRS:
        folder = Path(raw).expanduser()
        if folder.is_dir():
            installed.update(item.name.lower() for item in folder.iterdir())
    pairs = [(family, standin) for family, prefix, standin in _MAC_FONT_STANDINS
             if not any(name.startswith(prefix) for name in installed)]
    if not pairs:
        return
    entries = "".join(
        '<item oor:path="/org.openoffice.Office.Common/Font/Substitution/FontPairs">'
        f'<node oor:name="_{index}" oor:op="replace">'
        '<prop oor:name="Always" oor:op="fuse"><value>true</value></prop>'
        '<prop oor:name="OnScreenOnly" oor:op="fuse"><value>false</value></prop>'
        f'<prop oor:name="ReplaceFont" oor:op="fuse"><value>{family}</value></prop>'
        f'<prop oor:name="SubstituteFont" oor:op="fuse"><value>{standin}</value></prop>'
        "</node></item>\n"
        for index, (family, standin) in enumerate(pairs)
    )
    config = profile / "user" / "registrymodifications.xcu"
    config.parent.mkdir(parents=True, exist_ok=True)
    config.write_text(
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" '
        'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\n'
        '<item oor:path="/org.openoffice.Office.Common/Font/Substitution">'
        '<prop oor:name="Replacement" oor:op="fuse"><value>true</value></prop></item>\n'
        f"{entries}</oor:items>\n",
        encoding="utf-8",
    )


def _prefer_cjk_font(docx: Path) -> None:
    """Give a CJK document one installed font so the PDF does not mix fallbacks.

    East Asian text and the Latin text around it (digits, words, italics) share the
    font, so an italic number in a Chinese sentence does not switch to an unrelated
    serif. Fonts a style names explicitly, such as the code font, stay. Only the
    temporary copy converted to PDF is changed; the Word export keeps its theme fonts,
    which Word resolves itself.
    """
    if sys.platform != "darwin":
        return
    font = next((name for name, path in _MAC_CJK_FONTS if Path(path).is_file()), None)
    if font is None:
        return
    with zipfile.ZipFile(docx) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    contents = {info.filename: data for info, data in entries}
    if not _CJK_TEXT.search(contents.get("word/document.xml", b"").decode("utf-8", "ignore")):
        return
    styles = contents.get("word/styles.xml")
    if styles is None:
        return
    text = re.sub(r'\s+w:(?:ascii|hAnsi|eastAsia)Theme="[^"]*"', "", styles.decode("utf-8"))
    for attribute in ("eastAsia", "ascii", "hAnsi"):
        text = re.sub(rf'<w:rFonts\b(?![^>]*\bw:{attribute}=)', f'<w:rFonts w:{attribute}="{font}"', text)
    if not re.search(r"<w:rPrDefault>\s*<w:rPr>(?:(?!</w:rPr>)[\s\S])*<w:rFonts\b", text):
        text = re.sub(r"(<w:rPrDefault>\s*<w:rPr>)",
                      lambda found: f'{found.group(1)}<w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:eastAsia="{font}" />',
                      text, count=1)
    with zipfile.ZipFile(docx, "w", zipfile.ZIP_DEFLATED) as archive:
        for info, data in entries:
            archive.writestr(info, text.encode("utf-8") if info.filename == "word/styles.xml" else data)


def convert_to_pdf(soffice: str, docx: Path, output: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="astra-pdf-") as work:
        profile = Path(work) / "profile"
        _link_system_fonts(profile)
        _replace_missing_fonts(profile)
        source = Path(work) / f"{docx.stem or 'document'}.docx"
        shutil.copyfile(docx, source)
        _prefer_cjk_font(source)
        args = [
            soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--norestore",
            "--nolockcheck", "--convert-to", "pdf", "--outdir", work, str(source),
        ]
        completed = _run(args, timeout=240)
        produced = Path(work) / f"{source.stem}.pdf"
        if not produced.is_file():
            detail = (completed.stderr or completed.stdout or "").strip()[-2000:]
            raise ExportError("export_failed", f"LibreOffice could not convert the document to PDF: {detail}")
        _atomic_move(produced, output)


# ---- HTML (markdown-it-py) ---------------------------------------------------------------

_HTML_STYLE = """
body { margin: 0; background: #fff; color: #1f2328; font: 16px/1.6 -apple-system, "Segoe UI", "Noto Sans",
  "PingFang SC", "Microsoft YaHei", sans-serif; }
main { max-width: 780px; margin: 40px auto; padding: 0 20px; }
h1, h2, h3, h4 { line-height: 1.25; margin: 1.6em 0 0.6em; }
table { border-collapse: collapse; margin: 1em 0; width: 100%; }
th, td { border: 1px solid #d0d7de; padding: 6px 10px; text-align: left; vertical-align: top; }
th { background: #f6f8fa; }
pre { background: #f6f8fa; padding: 12px; overflow-x: auto; border-radius: 6px; }
code { font-family: ui-monospace, Menlo, Consolas, monospace; font-size: 0.9em; }
blockquote { margin: 1em 0; padding: 0 1em; color: #59636e; border-left: 4px solid #d0d7de; }
img { max-width: 100%; }
""".strip()


def export_html(markdown: str, *, output: Path, title: str | None, images: dict[str, Path]) -> None:
    from markdown_it import MarkdownIt

    parser = MarkdownIt("commonmark", {"html": False}).enable(["table", "strikethrough"])
    tokens = parser.parse(markdown)
    for token in tokens:
        for child in token.children or []:
            if child.type == "image":
                src = str(child.attrGet("src") or "")
                path = _lookup(images, src)
                if path is not None:
                    relative = os.path.relpath(path, output.parent)
                    child.attrSet("src", Path(relative).as_posix())
    body = parser.renderer.render(tokens, parser.options, {})
    page_title = _escape(title or output.stem)
    page = (
        f'<!doctype html>\n<html>\n<head>\n<meta charset="utf-8">\n'
        f'<meta name="viewport" content="width=device-width, initial-scale=1">\n'
        f"<title>{page_title}</title>\n<style>\n{_HTML_STYLE}\n</style>\n</head>\n"
        f"<body>\n<main>\n{body}</main>\n</body>\n</html>\n"
    )
    with tempfile.TemporaryDirectory(prefix="astra-export-") as work:
        target = Path(work) / "output.html"
        target.write_text(page, encoding="utf-8")
        _atomic_move(target, output)


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


# ---- Word (python-docx) ------------------------------------------------------------------


class _DocxBuilder:
    """Render markdown-it-py tokens into a python-docx document."""

    CODE_FONT = "Consolas"

    def __init__(self, *, template: Path | None, paper: str, images: dict[str, Path], warnings: list[str]):
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Mm

        self.qn = qn
        self.document = Document(str(template)) if template is not None else Document()
        if template is not None:
            body = self.document.element.body
            for element in list(body):
                if element.tag != qn("w:sectPr"):
                    body.remove(element)
        section = self.document.sections[0]
        if template is None:
            width, height = (Mm(210), Mm(297)) if paper == "a4" else (Mm(215.9), Mm(279.4))
            section.page_width, section.page_height = width, height
            for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
                setattr(section, side, Mm(25))
        page_width = section.page_width or Mm(210)
        margins = (section.left_margin or 0) + (section.right_margin or 0)
        self.text_width = int(page_width) - int(margins)
        self.images = images
        self.warnings = warnings
        self.lists: list[dict[str, Any]] = []
        self.quote_depth = 0
        self.title_mode = False
        self.used_title = False

    # -- style helpers --

    def _style(self, *names: str) -> str | None:
        """The first of ``names`` that the document (or its template) defines."""
        styles = self.document.styles
        for name in names:
            try:
                styles[name]
            except KeyError:
                continue
            return name
        return None

    def _paragraph(self, *styles: str):
        paragraph = self.document.add_paragraph()
        style = self._style(*styles) if styles else None
        if style is not None:
            paragraph.style = style
        return paragraph

    # -- rendering --

    def render(self, markdown: str, title: str | None) -> None:
        from markdown_it import MarkdownIt

        parser = MarkdownIt("commonmark", {"html": True}).enable(["table", "strikethrough"])
        tokens = parser.parse(markdown)
        headings = [token for token in tokens if token.type == "heading_open"]
        self.title_mode = bool(
            headings and headings[0].tag == "h1" and sum(token.tag == "h1" for token in headings) == 1
        )
        if title:
            self.document.core_properties.title = title
        index = 0
        while index < len(tokens):
            index = self._block(tokens, index)

    def _block(self, tokens: list[Any], index: int) -> int:
        token = tokens[index]
        kind = token.type
        if kind == "heading_open":
            self._heading(int(token.tag[1]), tokens[index + 1])
            return index + 3
        if kind == "paragraph_open":
            self._paragraph_block(tokens[index + 1])
            return index + 3
        if kind in ("bullet_list_open", "ordered_list_open"):
            start = token.attrGet("start")
            self.lists.append({"ordered": kind == "ordered_list_open", "counter": int(start or 1), "fresh": False})
            return index + 1
        if kind in ("bullet_list_close", "ordered_list_close"):
            self.lists.pop()
            return index + 1
        if kind == "list_item_open":
            if self.lists:
                self.lists[-1]["fresh"] = True
            return index + 1
        if kind == "list_item_close":
            if self.lists:
                self.lists[-1]["fresh"] = False
                if self.lists[-1]["ordered"]:
                    self.lists[-1]["counter"] += 1
            return index + 1
        if kind == "blockquote_open":
            self.quote_depth += 1
            return index + 1
        if kind == "blockquote_close":
            self.quote_depth -= 1
            return index + 1
        if kind in ("fence", "code_block"):
            self._code(token.content)
            return index + 1
        if kind == "table_open":
            return self._table(tokens, index)
        if kind == "hr":
            self._rule()
            return index + 1
        if kind == "html_block":
            if token.content.strip():
                self.warnings.append("Raw HTML block skipped in the Word export.")
            return index + 1
        return index + 1

    def _heading(self, level: int, inline: Any) -> None:
        if self.title_mode and level == 1 and not self.used_title:
            self.used_title = True
            style_names = ("Title",)
        else:
            mapped = max(1, level - 1) if self.title_mode else level
            style_names = (f"Heading {min(mapped, 9)}",)
        paragraph = self._paragraph(*style_names)
        if paragraph.style is None or paragraph.style.name not in style_names:
            self._inline(paragraph, inline.children or [], bold=True)
        else:
            self._inline(paragraph, inline.children or [])

    def _list_style(self) -> tuple[tuple[str, ...], str]:
        depth = len(self.lists)
        current = self.lists[-1]
        suffix = "" if depth <= 1 else f" {min(depth, 3)}"
        if current["ordered"]:
            return ("List Paragraph", "List"), f"{current['counter']}. "
        return (f"List Bullet{suffix}", "List Bullet", "List Paragraph"), ""

    def _paragraph_block(self, inline: Any) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Mm

        children = list(inline.children or [])
        meaningful = [child for child in children if child.type not in ("softbreak", "hardbreak")]
        if meaningful and all(child.type == "image" for child in meaningful):
            for child in meaningful:
                paragraph = self._paragraph("Normal")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._image(paragraph, child)
                alt = child.content or ""
                if alt:
                    caption = self._paragraph("Caption")
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = caption.add_run(alt)
                    run.italic = True
            return
        prefix = ""
        if self.lists and self.lists[-1]["fresh"]:
            box = self._task_box(children)
            ordered = self.lists[-1]["ordered"]
            # A task item shows its box where the bullet would be, not a bullet and a box.
            styles, prefix = (("List Paragraph", "List"), "") if box and not ordered else self._list_style()
            paragraph = self._paragraph(*styles)
            if paragraph.style is None or paragraph.style.name not in ("List Bullet", "List Bullet 2", "List Bullet 3"):
                if not ordered and not box:
                    prefix = "\u2022 "
                paragraph.paragraph_format.left_indent = Mm(6 * len(self.lists))
            elif len(self.lists) > 1:
                # Nested bullets indent past their parent item, whatever the parent list type.
                paragraph.paragraph_format.left_indent = Mm(6.35 * (len(self.lists) + 1))
            prefix += box
            self.lists[-1]["fresh"] = False
        elif self.lists:
            paragraph = self._paragraph("List Continue", "List Paragraph")
            paragraph.paragraph_format.left_indent = Mm(6 * len(self.lists))
        elif self.quote_depth:
            paragraph = self._paragraph("Quote")
        else:
            paragraph = self._paragraph()
        if prefix:
            paragraph.add_run(prefix)
        self._inline(paragraph, children)

    @staticmethod
    def _task_box(children: list[Any]) -> str:
        """Strip a list item's GitHub task marker and return the box that replaces it.

        The boxes are the ones pandoc uses; U+2611 can turn into a colour emoji when
        LibreOffice makes the PDF.
        """
        if not children or children[0].type != "text":
            return ""
        first = children[0]
        match = re.match(r"^\[([ xX])\]\s+", first.content)
        if not match:
            return ""
        first.content = first.content[match.end():]
        return "\u2612 " if match.group(1).lower() == "x" else "\u2610 "

    def _inline(self, paragraph: Any, children: list[Any], *, bold: bool = False) -> None:
        state = {"bold": 1 if bold else 0, "italic": 0, "strike": 0}
        link: dict[str, Any] | None = None
        for child in children:
            kind = child.type
            if kind == "text":
                if link is not None:
                    link["text"] += child.content
                else:
                    self._run(paragraph, child.content, state)
            elif kind == "softbreak":
                if link is not None:
                    link["text"] += " "
                else:
                    self._run(paragraph, " ", state)
            elif kind == "hardbreak":
                paragraph.add_run().add_break()
            elif kind in ("strong_open", "strong_close"):
                state["bold"] += 1 if kind.endswith("open") else -1
            elif kind in ("em_open", "em_close"):
                state["italic"] += 1 if kind.endswith("open") else -1
            elif kind in ("s_open", "s_close"):
                state["strike"] += 1 if kind.endswith("open") else -1
            elif kind == "code_inline":
                if link is not None:
                    link["text"] += child.content
                else:
                    run = self._run(paragraph, child.content, state)
                    run.font.name = self.CODE_FONT
            elif kind == "link_open":
                link = {"href": str(child.attrGet("href") or ""), "text": ""}
            elif kind == "link_close" and link is not None:
                self._link(paragraph, link["href"], link["text"], state)
                link = None
            elif kind == "image":
                self._image(paragraph, child)

    def _run(self, paragraph: Any, text: str, state: dict[str, int]) -> Any:
        run = paragraph.add_run(text)
        run.bold = bool(state["bold"]) or None
        run.italic = bool(state["italic"]) or None
        if state["strike"]:
            run.font.strike = True
        return run

    def _link(self, paragraph: Any, href: str, text: str, state: dict[str, int]) -> None:
        from docx.opc.constants import RELATIONSHIP_TYPE
        from docx.oxml import OxmlElement
        from docx.shared import RGBColor

        label = text or href
        if not re.match(r"^(https?|mailto):", href, re.IGNORECASE):
            run = self._run(paragraph, label, state)
            run.underline = True
            return
        qn = self.qn
        relationship = paragraph.part.relate_to(href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relationship)
        run = paragraph.add_run(label)
        run.bold = bool(state["bold"]) or None
        run.italic = bool(state["italic"]) or None
        run.underline = True
        run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
        style = self._style("Hyperlink")
        if style is not None:
            run.style = style
        hyperlink.append(run._r)
        paragraph._p.append(hyperlink)

    def _image(self, paragraph: Any, token: Any) -> None:
        from docx.shared import Emu

        src = _clean_src(str(token.attrGet("src") or ""))
        path = _lookup(self.images, src)
        if path is None:
            paragraph.add_run(f"[image: {token.content or src}]")
            return
        picture = path
        converted: Path | None = None
        try:
            from PIL import Image

            with Image.open(path) as image:
                width_px, _height = image.size
                dpi = float((image.info.get("dpi") or (96, 96))[0] or 96)
                if path.suffix.lower() in CONVERTIBLE_SUFFIXES:
                    handle = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
                    handle.close()
                    converted = Path(handle.name)
                    image.convert("RGBA").save(converted, "PNG")
                    picture = converted
            natural = int(width_px / max(dpi, 1) * 914400)
            width = Emu(min(natural, self.text_width))
            paragraph.add_run().add_picture(str(picture), width=width)
        except Exception as exc:  # corrupt or unsupported image: keep the export going
            self.warnings.append(f"Image could not be embedded: {src} ({exc})")
            paragraph.add_run(f"[image: {token.content or src}]")
        finally:
            if converted is not None:
                converted.unlink(missing_ok=True)

    def _code(self, content: str) -> None:
        from docx.oxml import OxmlElement
        from docx.shared import Pt

        paragraph = self._paragraph("No Spacing")
        properties = paragraph._p.get_or_add_pPr()
        shading = OxmlElement("w:shd")
        shading.set(self.qn("w:val"), "clear")
        shading.set(self.qn("w:color"), "auto")
        shading.set(self.qn("w:fill"), "F2F2F2")
        properties.append(shading)
        lines = content.rstrip("\n").split("\n")
        for position, line in enumerate(lines):
            run = paragraph.add_run(line)
            run.font.name = self.CODE_FONT
            run.font.size = Pt(9.5)
            if position < len(lines) - 1:
                run.add_break()
        self._paragraph()  # spacing after the block

    def _rule(self) -> None:
        from docx.oxml import OxmlElement

        paragraph = self._paragraph()
        properties = paragraph._p.get_or_add_pPr()
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "auto")):
            bottom.set(self.qn(key), value)
        border.append(bottom)
        properties.append(border)

    def _table(self, tokens: list[Any], index: int) -> int:
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        rows: list[list[tuple[Any, str, bool]]] = []
        position = index + 1
        while position < len(tokens) and tokens[position].type != "table_close":
            token = tokens[position]
            if token.type == "tr_open":
                rows.append([])
            elif token.type in ("th_open", "td_open"):
                style = str(token.attrGet("style") or "")
                align = style.split("text-align:")[-1].strip() if "text-align:" in style else ""
                rows[-1].append((tokens[position + 1], align, token.type == "th_open"))
            position += 1
        if rows:
            columns = max(len(row) for row in rows)
            table = self.document.add_table(rows=len(rows), cols=columns)
            style = self._style("Table Grid")
            if style is not None:
                table.style = style
            alignments = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
            for row_index, row in enumerate(rows):
                for column_index, (inline, align, header) in enumerate(row):
                    paragraph = table.cell(row_index, column_index).paragraphs[0]
                    if align in alignments:
                        paragraph.alignment = alignments[align]
                    self._inline(paragraph, list(inline.children or []), bold=header)
            self._paragraph()
        return position + 1

    def save(self, output: Path) -> None:
        with tempfile.TemporaryDirectory(prefix="astra-export-") as work:
            target = Path(work) / "output.docx"
            self.document.save(str(target))
            _atomic_move(target, output)


def export_with_python_docx(
    markdown: str,
    *,
    output: Path,
    title: str | None,
    template: Path | None,
    paper: str,
    images: dict[str, Path],
    warnings: list[str],
) -> None:
    builder = _DocxBuilder(template=template, paper=paper, images=images, warnings=warnings)
    builder.render(markdown, title)
    builder.save(output)


# ---- entry point -------------------------------------------------------------------------


def export_document(
    text: str,
    *,
    base_dir: Path,
    output: Path,
    fmt: str,
    engine: str = "auto",
    template: Path | None = None,
    paper: str = "a4",
    resolve_image: ImageResolver,
) -> ExportResult:
    """Export Markdown text; raises ExportError with a stable code on failure."""
    if fmt not in FORMATS:
        raise ExportError("invalid_format", f"Unsupported export format: {fmt}")
    if engine not in ENGINES:
        raise ExportError("invalid_engine", f"Unsupported export engine: {engine}")
    warnings: list[str] = []
    markdown, pending, title = prepare_markdown(text)
    if pending:
        warnings.append(f"{pending} section(s) are still pending placeholders.")
    markdown, images = check_images(markdown, resolve_image, warnings)

    pandoc = find_pandoc() if engine in ("auto", "pandoc") else None
    if engine == "pandoc" and pandoc is None:
        raise ExportError("export_unavailable", "pandoc is not installed.", "Install pandoc or use engine python-docx.")
    use_pandoc = pandoc is not None and (engine == "pandoc" or engine == "auto")
    if not use_pandoc:
        needed = markdown_it_available() if fmt == "html" else python_docx_available()
        if not needed:
            raise ExportError("export_unavailable", "The Word/HTML exporter is not installed.", INSTALL_HINT)
    chosen = "pandoc" if use_pandoc else ("markdown-it" if fmt == "html" else "python-docx")

    soffice = None
    if fmt == "pdf":
        soffice = find_soffice()
        if soffice is None:
            raise ExportError(
                "export_unavailable",
                "PDF export needs LibreOffice, which was not found.",
                "Install LibreOffice (or set ASTRA_SOFFICE), or export docx and print it to PDF.",
            )

    with tempfile.TemporaryDirectory(prefix="astra-export-") as work:
        stage = Path(work) / f"{output.stem or 'document'}.docx" if fmt == "pdf" else output
        if use_pandoc:
            assert pandoc is not None
            export_with_pandoc(
                pandoc, markdown, base_dir=base_dir, output=stage, fmt="docx" if fmt == "pdf" else fmt,
                title=title, template=template, warnings=warnings, paper=paper,
            )
        elif fmt == "html":
            export_html(markdown, output=stage, title=title, images=images)
        else:
            export_with_python_docx(
                markdown, output=stage, title=title, template=template, paper=paper,
                images=images, warnings=warnings,
            )
        if fmt == "pdf":
            assert soffice is not None
            convert_to_pdf(soffice, stage, output)
            chosen += " + LibreOffice"
    return ExportResult(output=output, engine=chosen, warnings=warnings)
