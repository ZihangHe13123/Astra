"""The canvas tool (specs/canvas.md): it makes a deck, hands the model the template's guide, and asks the
desktop for a slide's picture or a print of the deck."""

from __future__ import annotations

import asyncio
import hashlib
import json
from pathlib import Path

import pytest

from agent.runtime.canvas import slide_count
from agent.runtime.page_snapshot import PageSnapshot, PageSnapshotBroker, PageSnapshotError
from agent.runtime.tools.canvas import register_canvas_tool
from agent.runtime.tools.registry import ToolRegistry


class Desktop:
    """Answers render requests as the desktop would, and remembers what it was asked."""

    def __init__(self, tmp_path: Path, error: str = ""):
        self.asked: list[tuple[Path, dict]] = []
        self.tmp_path = tmp_path
        self.error = error

    async def __call__(self, path: Path, **options) -> PageSnapshot:
        self.asked.append((path, options))
        if self.error:
            raise PageSnapshotError(self.error)
        if options.get("export") == "slides":
            from PIL import Image

            folder = Path(str(options["target"]))
            folder.mkdir(parents=True, exist_ok=True)
            count = slide_count(path.parent)
            pictures = []
            for number in range(1, count + 1):
                picture = folder / f"slide-{number:03d}.png"
                Image.new("RGB", (192, 108), (number * 40, 120, 80)).save(picture)
                pictures.append(picture)
            return PageSnapshot(image_path=folder, paths=tuple(pictures))
        if options.get("export") == "pdf":
            target = Path(str(options["target"]))
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(b"%PDF-1.7 fixture")
            return PageSnapshot(image_path=target)
        picture = self.tmp_path / f"slide-{options.get('slide')}.png"
        picture.write_bytes(b"\x89PNG fixture")
        return PageSnapshot(image_path=picture, height=720, page_height=1080)


def tool(tmp_path: Path, desktop: Desktop):
    registry = ToolRegistry()
    register_canvas_tool(registry, workdir=tmp_path, render=desktop)

    def call(**args):
        # Looking is a tool of its own (read, no approval); the rest is canvas.
        if args.get("action") == "look":
            args.pop("action")
            return asyncio.run(registry.execute("canvas_look", args))
        return asyncio.run(registry.execute("canvas", args))
    return call


def test_looking_is_a_read_of_its_own_and_making_or_printing_is_a_write(tmp_path):
    registry = ToolRegistry()
    register_canvas_tool(registry, workdir=tmp_path, render=Desktop(tmp_path))
    canvas, look = registry.get("canvas"), registry.get("canvas_look")
    assert canvas is not None and look is not None
    assert (canvas.risk, look.risk) == ("write", "read")
    assert canvas.parameters["properties"]["action"]["enum"] == ["types", "create", "export"]


def test_create_makes_a_deck_in_the_workspace_and_hands_over_the_guide(tmp_path):
    call = tool(tmp_path, Desktop(tmp_path))
    result = call(action="create", title="季度回顾 Q3")
    assert not result["error"]
    folder = tmp_path / "季度回顾-Q3.canvas"
    assert result["output"].splitlines()[0] == f"Canvas: {folder / 'canvas.json'}"
    assert "# Slides canvas" in result["output"] and str(folder / "slides.html") in result["output"]
    manifest = json.loads((folder / "canvas.json").read_text(encoding="utf-8"))
    assert (manifest["type"], manifest["title"], manifest["entry"]) == ("slides", "季度回顾 Q3", "slides.html")
    assert (folder / "assets").is_dir() and slide_count(folder) == 1
    # Replacing the starter slide whole needs its hash, which write_file asks for; the result gives it.
    starter = hashlib.sha256((folder / "slides.html").read_bytes()).hexdigest()
    assert f'write_file(overwrite=true, expected_sha256="{starter}")' in result["output"]
    # The same title again is a second deck; a named path is used as given; nothing is made outside the workspace.
    assert call(action="create", title="季度回顾 Q3")["output"].startswith(f"Canvas: {tmp_path / '季度回顾-Q3-2.canvas'}")
    assert call(action="create", title="x", path="decks/plan")["output"].startswith(f"Canvas: {tmp_path / 'decks/plan.canvas'}")
    assert "already exists" in call(action="create", title="x", path="decks/plan.canvas")["error"]
    assert "inside the workspace" in call(action="create", title="x", path=str(tmp_path.parent / "out"))["error"]
    assert "Unknown canvas type" in call(action="create", type="dashboard", title="x")["error"]


def test_look_asks_the_desktop_for_one_slide_and_attaches_its_picture(tmp_path):
    desktop = Desktop(tmp_path)
    call = tool(tmp_path, desktop)
    call(action="create", title="deck")
    (tmp_path / "deck.canvas" / "slides.html").write_text(
        '<section class="slide title"><h1>A</h1></section>\n<section class="big slide"><p>B</p></section>', encoding="utf-8")
    result = call(action="look", canvas="deck.canvas", slide=2)
    payload = json.loads(result["output"])
    assert (payload["type"], payload["slide"], payload["slides"]) == ("image_attachment", 2, 2)
    assert payload["image_paths"] == [str(tmp_path / "slide-2.png")]
    assert payload["canvas"] == str(tmp_path / "deck.canvas" / "canvas.json")
    # The desktop is asked by the canvas's canvas.json, which is what it opens; a file inside the canvas names it too.
    assert desktop.asked == [(tmp_path / "deck.canvas" / "canvas.json", {"slide": 2})]
    assert not call(action="look", canvas=str(tmp_path / "deck.canvas" / "slides.html"), slide=1)["error"]
    assert "slide is 1 to 2" in call(action="look", canvas="deck.canvas", slide=3)["error"]
    assert "not in a canvas" in call(action="look", canvas="elsewhere", slide=1)["error"]


def test_export_prints_into_the_canvas_exports_folder(tmp_path):
    desktop = Desktop(tmp_path)
    call = tool(tmp_path, desktop)
    call(action="create", title="deck")
    result = json.loads(call(action="export", canvas="deck.canvas")["output"])
    target = tmp_path / "deck.canvas" / "exports" / "deck.pdf"
    assert result == {"success": True, "file": str(target), "format": "pdf", "slides": 1,
                      "canvas": str(tmp_path / "deck.canvas" / "canvas.json"), "title": "deck"}
    assert desktop.asked[-1] == (tmp_path / "deck.canvas" / "canvas.json", {"export": "pdf", "target": str(target), "timeout": 150})
    assert "pdf or pptx" in call(action="export", canvas="deck.canvas", format="docx")["error"]


def test_a_desktop_that_cannot_draw_is_reported_once_and_not_retried(tmp_path):
    call = tool(tmp_path, Desktop(tmp_path, error="页面没有载入。"))
    call(action="create", title="deck")
    result = call(action="look", canvas="deck.canvas", slide=1)
    assert (result["error"], result["code"], result["retryable"]) == ("页面没有载入。", "canvas_rendering_failed", False)


def test_the_broker_carries_options_and_takes_a_written_file_as_the_answer():
    sent: list[dict] = []
    broker = PageSnapshotBroker(sent.append)

    async def scenario():
        asked = asyncio.ensure_future(broker.request(Path("/w/deck.canvas/canvas.json"), export="pdf", target="/w/x.pdf", timeout=5))
        await asyncio.sleep(0)
        assert sent[0]["export"] == "pdf" and sent[0]["target"] == "/w/x.pdf" and "timeout" not in sent[0]
        assert broker.resolve({"request_id": sent[0]["request_id"], "file_path": "/w/x.pdf"})
        return await asked

    assert asyncio.run(scenario()).image_path == Path("/w/x.pdf")


@pytest.mark.parametrize("markup, count", [
    ('<section class="slide">a</section><section class="slide title">b</section>', 2),
    ("<section class='intro slide wide'>a</section><section class=\"slides\">no</section><div class=\"slide\">no</div>", 1),
])
def test_slides_are_counted_as_the_desktop_counts_them(tmp_path, markup, count):
    folder = tmp_path / "x.canvas"
    folder.mkdir()
    (folder / "slides.html").write_text(markup, encoding="utf-8")
    assert slide_count(folder) == count


def test_a_pptx_is_one_picture_per_slide_with_its_heading_and_notes(tmp_path):
    from pptx import Presentation

    desktop = Desktop(tmp_path)
    call = tool(tmp_path, desktop)
    call(action="create", title="季度回顾")
    (tmp_path / "季度回顾.canvas" / "slides.html").write_text(
        '<section class="slide title"><h1>季度 &amp; 回顾</h1><aside class="notes">先讲<b>营收</b>。</aside></section>\n'
        '<section class="slide"><h2>没有备注</h2><p>x</p></section>', encoding="utf-8")
    result = json.loads(call(action="export", canvas="季度回顾.canvas", format="pptx")["output"])
    target = tmp_path / "季度回顾.canvas" / "exports" / "季度回顾.pptx"
    assert (result["file"], result["format"], result["slides"]) == (str(target), "pptx", 2)
    assert desktop.asked[-1][1]["export"] == "slides"
    # The pictures were only for the PPTX.
    assert not (tmp_path / "季度回顾.canvas" / "exports" / ".slides").exists()
    deck = Presentation(str(target))
    assert (deck.slide_width, deck.slide_height) == (12192000, 6858000)
    assert len(deck.slides) == 2
    pictures = [slide.shapes[0] for slide in deck.slides]
    assert [(shape.left, shape.top, shape.width, shape.height) for shape in pictures] == [(0, 0, 12192000, 6858000)] * 2
    assert [shape._element.nvPicPr.cNvPr.get("descr") for shape in pictures] == ["季度 & 回顾", "没有备注"]
    assert deck.slides[0].notes_slide.notes_text_frame.text == "先讲 营收 。"
    assert not deck.slides[1].has_notes_slide
    assert deck.core_properties.title == "季度回顾"
