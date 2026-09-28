import asyncio
import json
import os
import sys
import zipfile
from pathlib import Path

import pytest

from agent.runtime.tools import documents_export
from agent.runtime.tools.documents import register_document_tools
from agent.runtime.tools.documents_export import (
    ExportError,
    check_images,
    export_document,
    prepare_markdown,
)
from agent.runtime.tools.files import register_file_tools
from agent.runtime.tools.registry import ToolRegistry

needs_docx = pytest.mark.skipif(
    not documents_export.python_docx_available(), reason="documents extra (python-docx) not installed"
)
needs_markdown_it = pytest.mark.skipif(
    not documents_export.markdown_it_available(), reason="markdown-it-py not installed"
)


def run(coro):
    return asyncio.run(coro)


def png(path: Path, size: tuple[int, int] = (320, 160)) -> Path:
    from PIL import Image

    Image.new("RGB", size, (30, 120, 200)).save(path, "PNG")
    return path


def inside(base: Path):
    def resolve(src: str) -> Path:
        candidate = Path(src)
        if not candidate.is_absolute():
            candidate = base / candidate
        resolved = candidate.resolve()
        if base.resolve() not in (resolved, *resolved.parents):
            raise ValueError("outside the workspace")
        return resolved
    return resolve


SAMPLE = """# Proposal

The plan in one line.

<!-- astra:section id="goals" -->
## Goals

1. First goal with **bold** and `code`.
2. Second goal with a [link](https://example.com/a).
    - nested bullet
- [x] done item
- [ ] open item

<!-- astra:section id="data" status="pending" intent="Participants and labels" -->
## Data

*Pending: Participants and labels*

<!-- astra:section id="results" -->
## Results

| Method | F1 |
| --- | ---: |
| Rules | 0.80 |
| CNN | 0.91 |

![Rotation chart](figures/chart.png)

> A quoted remark. <!-- @astra: soften this -->

```python
# not a heading
print("hi")
```

---

Final paragraph.
"""


def test_prepare_markdown_strips_markers_and_comments_but_keeps_placeholders():
    markdown, pending, title = prepare_markdown(SAMPLE)
    assert title == "Proposal"
    assert pending == 1
    assert "astra:section" not in markdown and "@astra" not in markdown
    assert "*Pending: Participants and labels*" in markdown
    assert "# not a heading" in markdown


def test_check_images_keeps_local_files_and_replaces_the_rest(tmp_path):
    png(tmp_path / "ok.png")
    (tmp_path / "vector.svg").write_text("<svg/>", encoding="utf-8")
    warnings: list[str] = []
    markdown, images = check_images(
        "![ok](ok.png)\n![remote](https://example.com/x.png)\n![gone](missing.png)\n"
        "![vector](vector.svg)\n![escape](../../etc/passwd.png)\n```\n![code](code.png)\n```\n",
        inside(tmp_path),
        warnings,
    )
    assert list(images) == ["ok.png"]
    assert "![ok](ok.png)" in markdown
    assert "[image: remote]" in markdown and "[image: gone]" in markdown
    assert "[image: vector]" in markdown and "[image: escape]" in markdown
    assert "![code](code.png)" in markdown  # fenced code is left alone
    assert len(warnings) == 4


@needs_docx
def test_python_docx_export_renders_structure(tmp_path, monkeypatch):
    from docx import Document

    monkeypatch.setattr(documents_export, "find_pandoc", lambda: None)
    (tmp_path / "figures").mkdir()
    png(tmp_path / "figures" / "chart.png")
    output = tmp_path / "proposal.docx"
    result = export_document(
        SAMPLE, base_dir=tmp_path, output=output, fmt="docx", resolve_image=inside(tmp_path),
    )
    assert result.engine == "python-docx"
    assert any("pending" in warning for warning in result.warnings)

    document = Document(str(output))
    paragraphs = [(item.style.name, item.text) for item in document.paragraphs]
    assert ("Title", "Proposal") in paragraphs
    assert ("Heading 1", "Goals") in paragraphs and ("Heading 1", "Results") in paragraphs
    texts = [text for _, text in paragraphs]
    assert "1. First goal with bold and code." in texts
    assert any(text.startswith("2. Second goal with a link") for text in texts)
    assert "☑ done item" in texts and "☐ open item" in texts
    assert "A quoted remark." in texts
    assert not any("@astra" in text or "astra:section" in text for text in texts)
    table = document.tables[0]
    assert [cell.text for cell in table.rows[0].cells] == ["Method", "F1"]
    assert table.rows[0].cells[0].paragraphs[0].runs[0].bold
    assert len(document.inline_shapes) == 1
    assert any("hyperlink" in rel.reltype for rel in document.part.rels.values())
    code = next(item for item in document.paragraphs if "print(" in item.text)
    assert code.runs[0].font.name == "Consolas"
    assert document.core_properties.title == "Proposal"
    width = document.sections[0].page_width
    assert width is not None and round(width.mm) == 210


@needs_docx
def test_python_docx_export_reuses_template_styles(tmp_path, monkeypatch):
    from docx import Document
    from docx.shared import Pt

    monkeypatch.setattr(documents_export, "find_pandoc", lambda: None)
    template = Document()
    template.styles["Normal"].font.name = "Times New Roman"
    template.styles["Normal"].font.size = Pt(13)
    template.add_paragraph("template body is dropped")
    template.save(str(tmp_path / "template.docx"))
    output = tmp_path / "out.docx"
    export_document(
        "# T\n\nBody text.\n", base_dir=tmp_path, output=output, fmt="docx",
        template=tmp_path / "template.docx", resolve_image=inside(tmp_path),
    )
    document = Document(str(output))
    assert document.styles["Normal"].font.name == "Times New Roman"
    assert [item.text for item in document.paragraphs] == ["T", "Body text."]


@needs_markdown_it
def test_html_export_links_images_relative_to_the_output(tmp_path, monkeypatch):
    monkeypatch.setattr(documents_export, "find_pandoc", lambda: None)
    (tmp_path / "figures").mkdir()
    png(tmp_path / "figures" / "chart.png")
    output = tmp_path / "out" / "page.html"
    result = export_document(SAMPLE, base_dir=tmp_path, output=output, fmt="html", resolve_image=inside(tmp_path))
    html = output.read_text(encoding="utf-8")
    assert result.engine == "markdown-it"
    assert "<title>Proposal</title>" in html and "<table>" in html
    assert 'src="../figures/chart.png"' in html
    assert "@astra" not in html and "astra:section" not in html


def fake_pandoc(tmp_path: Path) -> Path:
    script = tmp_path / "fake_pandoc.py"
    log = tmp_path / "pandoc-args.json"
    script.write_text(
        "import json, shutil, sys\n"
        "args = sys.argv[1:]\n"
        f"json.dump({{'args': args, 'input': open(args[0], encoding='utf-8').read()}}, open({str(log)!r}, 'w'))\n"
        "open(args[args.index('--output') + 1], 'wb').write(b'fake docx')\n",
        encoding="utf-8",
    )
    if os.name == "nt":
        wrapper = tmp_path / "fake_pandoc.cmd"
        wrapper.write_text(f'@"{sys.executable}" "{script}" %*\n', encoding="utf-8")
        return wrapper
    wrapper = tmp_path / "fake_pandoc"
    wrapper.write_text(f'#!/bin/sh\nexec "{sys.executable}" "{script}" "$@"\n', encoding="utf-8")
    wrapper.chmod(0o755)
    return wrapper


def test_auto_engine_prefers_pandoc_with_safe_arguments(tmp_path, monkeypatch):
    monkeypatch.setenv("ASTRA_PANDOC", str(fake_pandoc(tmp_path)))
    template = tmp_path / "template.docx"
    template.write_bytes(b"template")
    output = tmp_path / "out.docx"
    result = export_document(
        "# T\n\n![remote](https://example.com/a.png)\n\nText <!-- @astra: note -->\n",
        base_dir=tmp_path, output=output, fmt="docx", template=template, resolve_image=inside(tmp_path),
    )
    assert result.engine == "pandoc"
    assert output.read_bytes() == b"fake docx"
    record = json.loads((tmp_path / "pandoc-args.json").read_text(encoding="utf-8"))
    args = record["args"]
    assert args[1:5] == ["--from", "gfm", "--to", "docx"]
    assert f"--resource-path={tmp_path}" in args
    assert args[args.index("--reference-doc") + 1] == str(template)
    assert "https://example.com" not in record["input"] and "[image: remote]" in record["input"]
    assert "@astra" not in record["input"]


def test_missing_engines_report_how_to_install(tmp_path, monkeypatch):
    monkeypatch.setattr(documents_export, "find_pandoc", lambda: None)
    monkeypatch.setattr(documents_export, "python_docx_available", lambda: False)
    with pytest.raises(ExportError) as missing:
        export_document("# T\n", base_dir=tmp_path, output=tmp_path / "t.docx", fmt="docx", resolve_image=inside(tmp_path))
    assert missing.value.code == "export_unavailable"
    assert "--extra documents" in missing.value.hint
    with pytest.raises(ExportError) as no_pandoc:
        export_document(
            "# T\n", base_dir=tmp_path, output=tmp_path / "t.docx", fmt="docx",
            engine="pandoc", resolve_image=inside(tmp_path),
        )
    assert no_pandoc.value.code == "export_unavailable"
    monkeypatch.setattr(documents_export, "python_docx_available", lambda: True)
    monkeypatch.setattr(documents_export, "find_soffice", lambda: None)
    with pytest.raises(ExportError) as no_office:
        export_document("# T\n", base_dir=tmp_path, output=tmp_path / "t.pdf", fmt="pdf", resolve_image=inside(tmp_path))
    assert no_office.value.code == "export_unavailable" and "LibreOffice" in str(no_office.value)


@needs_docx
@pytest.mark.skipif(documents_export.find_soffice() is None, reason="LibreOffice not installed")
def test_pdf_export_through_libreoffice(tmp_path, monkeypatch):
    pypdf = pytest.importorskip("pypdf")
    monkeypatch.setattr(documents_export, "find_pandoc", lambda: None)
    (tmp_path / "figures").mkdir()
    png(tmp_path / "figures" / "chart.png")
    output = tmp_path / "proposal.pdf"
    result = export_document(SAMPLE, base_dir=tmp_path, output=output, fmt="pdf", resolve_image=inside(tmp_path))
    assert result.engine == "python-docx + LibreOffice"
    assert output.read_bytes().startswith(b"%PDF")
    reader = pypdf.PdfReader(str(output))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "Proposal" in text and "Results" in text


# ---- the doc_export tool ---------------------------------------------------------------------


def make_registry(workspace: Path) -> ToolRegistry:
    registry = ToolRegistry()
    register_file_tools(registry, workdir=str(workspace))
    register_document_tools(registry)
    return registry


def test_export_tool_validates_output_and_needs_approval_outside_the_workspace(tmp_path, monkeypatch):
    monkeypatch.setenv("ASTRA_PANDOC", str(fake_pandoc(tmp_path)))
    workspace = tmp_path / "ws"
    workspace.mkdir()
    (workspace / "plan.md").write_text(
        '# Plan\n\n<!-- astra:section id="a" status="pending" intent="Later" -->\n## A\n\n*Pending: Later*\n',
        encoding="utf-8",
    )
    registry = make_registry(workspace)

    exported = json.loads(run(registry.execute("doc_export", {"path": "plan.md"}))["output"])
    assert exported["path"] == str(workspace / "plan.docx")
    assert exported["engine"] == "pandoc" and exported["pending_sections"] == 1
    assert exported["checkpoint_id"] and exported["bytes"] == len(b"fake docx")
    assert any("pending" in warning for warning in exported["warnings"])

    wrong = run(registry.execute("doc_export", {"path": "plan.md", "format": "pdf", "output": "plan.docx"}))
    assert wrong["code"] == "invalid_output"
    outside = run(registry.execute("doc_export", {"path": "plan.md", "output": str(tmp_path / "elsewhere.docx")}))
    assert "ToolApprovalRequired" in outside["error"]
    bad_format = run(registry.execute("doc_export", {"path": "plan.md", "format": "odt"}))
    assert bad_format["error"]


@needs_docx
def test_pdf_copy_pins_an_installed_cjk_font_on_macos(tmp_path, monkeypatch):
    from docx import Document

    font_file = tmp_path / "cjk.ttc"
    font_file.write_bytes(b"font")
    monkeypatch.setattr(documents_export.sys, "platform", "darwin")
    monkeypatch.setattr(documents_export, "_MAC_CJK_FONTS", (("Test CJK", str(font_file)),))
    chinese, english = tmp_path / "zh.docx", tmp_path / "en.docx"
    for path, text in ((chinese, "中文段落"), (english, "English only")):
        document = Document()
        document.add_paragraph(text)
        document.save(str(path))
    documents_export._prefer_cjk_font(chinese)
    documents_export._prefer_cjk_font(english)
    styles = zipfile.ZipFile(chinese).read("word/styles.xml").decode("utf-8")
    assert "eastAsiaTheme" not in styles and 'w:eastAsia="Test CJK"' in styles
    assert "eastAsiaTheme" in zipfile.ZipFile(english).read("word/styles.xml").decode("utf-8")
    assert Document(str(chinese)).paragraphs[0].text == "中文段落"


@pytest.mark.skipif(os.name == "nt", reason="symlinks need extra privileges on Windows")
def test_system_fonts_are_linked_into_the_libreoffice_profile_on_macos(tmp_path, monkeypatch):
    fonts = tmp_path / "fonts"
    fonts.mkdir()
    monkeypatch.setattr(documents_export, "_MAC_FONT_DIRS", (str(fonts), str(tmp_path / "missing")))
    monkeypatch.setattr(documents_export.sys, "platform", "linux")
    documents_export._link_system_fonts(tmp_path / "linux-profile")
    assert not (tmp_path / "linux-profile").exists()
    monkeypatch.setattr(documents_export.sys, "platform", "darwin")
    documents_export._link_system_fonts(tmp_path / "profile")
    linked = list((tmp_path / "profile" / "user" / "fonts").iterdir())
    assert [item.resolve() for item in linked] == [fonts.resolve()]


@needs_docx
@pytest.mark.skipif(sys.platform != "darwin" or documents_export.find_soffice() is None, reason="macOS LibreOffice check")
def test_chinese_text_is_drawn_with_a_cjk_font_in_the_pdf(tmp_path, monkeypatch):
    pypdf = pytest.importorskip("pypdf")
    monkeypatch.setattr(documents_export, "find_pandoc", lambda: None)
    output = tmp_path / "zh.pdf"
    export_document("# 手势识别评测\n\n侧对时规则明显下降。\n", base_dir=tmp_path, output=output, fmt="pdf", resolve_image=inside(tmp_path))
    fonts = set()
    for page in pypdf.PdfReader(str(output)).pages:
        for font in page["/Resources"]["/Font"].values():
            fonts.add(str(font.get_object()["/BaseFont"]))
    assert any(name in font for font in fonts for name in ("Hiragino", "STHeiti", "Songti", "ArialUnicode")), fonts
